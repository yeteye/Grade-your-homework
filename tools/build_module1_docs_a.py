"""Build member A Word deliverables from the versioned Markdown evidence."""
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'deliverables/module1'
OUT.mkdir(parents=True, exist_ok=True)


def set_font(style, name='Microsoft YaHei', size=10.5, bold=False):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.insert(0, fonts)
    fonts.set(qn('w:eastAsia'), name)


def configure(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.2)
    for name, size, bold in [('Normal', 10.5, False), ('Title', 17, True),
                             ('Heading 1', 13, True), ('Heading 2', 11, True)]:
        set_font(doc.styles[name], size=size, bold=bold)
    normal = doc.styles['Normal'].paragraph_format
    normal.space_after = Pt(6)
    normal.line_spacing = 1.25
    for name in ('Heading 1', 'Heading 2'):
        fmt = doc.styles[name].paragraph_format
        fmt.space_before = Pt(12)
        fmt.space_after = Pt(5)
        fmt.keep_with_next = True
    doc.styles['Title'].paragraph_format.space_after = Pt(10)
    title_ppr = doc.styles['Title'].element.get_or_add_pPr()
    border = title_ppr.find(qn('w:pBdr'))
    if border is not None:
        title_ppr.remove(border)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('软件测试与质量保证实践  |  曹浩 M202677234').font.size = Pt(8)


def add_table(doc, pairs):
    table = doc.add_table(rows=0, cols=2)
    table.autofit = False
    table.columns[0].width = Cm(3.5)
    table.columns[1].width = Cm(12.5)
    for index, (key, value) in enumerate(pairs):
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = value
        for cell in cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for side in ('top', 'left', 'bottom', 'right'):
                edge = OxmlElement(f'w:{side}')
                edge.set(qn('w:val'), 'single')
                edge.set(qn('w:sz'), '4')
                edge.set(qn('w:color'), 'D9D9D9')
                borders.append(edge)
            tcpr.append(borders)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                if index < len(pairs) - 1:
                    paragraph.paragraph_format.keep_with_next = True
                for run in paragraph.runs:
                    run.font.size = Pt(9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def render_markdown(doc, md, shift=0):
    for line in md.splitlines():
        line = line.strip()
        if not line:
            continue
        if line.startswith('#'):
            level = len(line) - len(line.lstrip('#'))
            text = line[level:].strip()
            if level == 1 and shift == 0:
                doc.add_paragraph(text, style='Title')
            else:
                doc.add_heading(text, level=min(2, max(1, level - 1 + shift)))
        elif re.match(r'^\d+\. ', line):
            text = re.sub(r'^\d+\.\s*', '', line)
            doc.add_paragraph(re.sub(r'`([^`]+)`', r'\1', text), style='List Number')
        else:
            doc.add_paragraph(re.sub(r'`([^`]+)`', r'\1', line))


def defect_document():
    doc = Document()
    configure(doc)
    doc.add_paragraph('阅知 Homework Studio 模块一缺陷报告', style='Title')
    doc.add_paragraph('曹浩  M202677234     2026-09-15     被测基线 2da8824')
    doc.add_paragraph('本报告记录成员 A 在评分与输入校验范围内复现并完成回归的四项缺陷。修复前和修复后结果均保存在仓库 reports/module1/；实际发现数量与修复结论以脚本和日志为准。')
    doc.add_heading('1 引言', level=1)
    doc.add_paragraph('对象为本地作业批改系统的单份评分流程，涉及请求参数、要点、Transformer 返回值和可选云端评分。本文只覆盖从 2da8824 开始的模块一测试活动。')
    doc.add_heading('2 测试环境', level=1)
    add_table(doc, [('平台', 'Windows x64，Python 3.12.14'),
                    ('支撑软件', 'Flask 3.1.3，Pillow 12.3.0，Werkzeug 3.1.8，SQLite'),
                    ('隔离方式', 'Flask test client、逐用例临时 SQLite 目录；模型和云端响应使用确定性替身'),
                    ('执行证据', 'a-defects-baseline.json：4 项失败；a-final-full.json：4 项通过，全部 43 项检查通过')])
    doc.add_heading('3 测试策略与执行步骤', level=1)
    doc.add_paragraph('对冻结基线运行 A-019～A-022，保存失败响应及记录数；逐项修复后重新执行，并复跑成员 A 的其余正式用例和原有开发检查。运行命令为 run-module1.bat --suite all --label <唯一名称>。')
    doc.add_heading('4 缺陷表', level=1)
    summaries = [('A-DEF-001','超大整数返回500','中','A-019','358a085'),
                 ('A-DEF-002','等价要点重复计权','中','A-020','ec22c91'),
                 ('A-DEF-003','非法模型概率生成有效成绩','高','A-021','103d09e'),
                 ('A-DEF-004','超大云端分数未回退','中','A-022','7ca7fc1')]
    for ident,title,severity,case,fix in summaries:
        doc.add_heading(f'{ident} {title}', level=2)
        add_table(doc, [('测试人', '曹浩 M202677234'), ('测试项', '评分与输入校验'),
                        ('用例编号', case), ('严重程度', severity), ('优先级', severity),
                        ('状态', '已修复并回归通过'), ('修复提交', fix)])
        raw = (ROOT / 'docs/module1/defects' / f'{ident}.md').read_text(encoding='utf-8')
        for section in raw.split('## ')[1:]:
            heading, body = section.split('\n', 1)
            doc.add_paragraph(heading.strip(), style='Heading 2')
            for paragraph in body.split('\n\n'):
                paragraph = paragraph.strip().replace('\n', ' ')
                if paragraph:
                    doc.add_paragraph(paragraph.replace('`', ''))
    doc.add_heading('5 结果分析和结论', level=1)
    doc.add_paragraph('四项缺陷均在 2da8824 上复现为失败用例，经过各自修复提交后全部通过。最终完整运行的 22 条成员 A 正式用例与 21 条原有开发检查均通过。真实模型质量和云端连接不在本次验证范围，最终小组质量结论待另一成员结果合并。')
    target = OUT / '曹浩-缺陷报告.docx'
    doc.save(target)
    return target


def report_document():
    doc = Document()
    configure(doc)
    doc.styles['Normal'].font.size = Pt(10)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.12
    doc.styles['Normal'].paragraph_format.space_after = Pt(4)
    md = (ROOT / 'docs/module1/report-a.md').read_text(encoding='utf-8')
    render_markdown(doc, md)
    target = OUT / '曹浩-测试报告章节.docx'
    doc.save(target)
    return target


if __name__ == '__main__':
    print(defect_document())
    print(report_document())
